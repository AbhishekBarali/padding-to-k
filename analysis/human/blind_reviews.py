"""Post hoc blind check of the reviews set (not in the registered plan).

The registered blind check left the reviews set out: its text may not be shared, so its documents went on a
separate local sheet that was never filled. This draws 5 fresh review documents (full split, never shown
with their answer list in a cleanup review page) with a fixed seed, writes one Excel sheet per author
WITHOUT the answer lists, and scores the filled sheets against the lists.

    uv run --no-project --with openpyxl==3.1.5 python analysis/human/blind_reviews.py make --out <dir>
    uv run --no-project --with openpyxl==3.1.5 python analysis/human/blind_reviews.py score A.xlsx B.xlsx --out analysis/human/blind_reviews_result.csv

The filled sheets of 20 Oct 2026 are in data/local/blind-reviews-{A,B}-filled.xlsx (not committed).

The sheets hold review text, so they and the key stay out of git (data/local is ignored).
"""
import argparse
import csv
import json
import random
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[1]
sys.path.insert(0, str(REPO / "data"))
sys.path.insert(0, str(HERE))
import blind_check  # noqa: E402
import forms  # noqa: E402

SEED = 20261020
N_DOCS = 5
KEY = REPO / "data" / "local" / "blind_key_reviews_posthoc.csv"
UNCLEAR = [("y", "Yes"), ("n", "No")]
COLS = [{"head": "No.", "width": 6, "kind": "no"},
        {"head": "Text", "width": 70, "kind": "read"},
        {"head": "What to list", "width": 30, "kind": "read"},
        {"head": "Items you found, one per line", "width": 34, "kind": "text"},
        {"head": "How many items?", "width": 12, "kind": "number"},
        {"head": "Could two careful readers count differently?", "width": 16, "kind": "select",
         "options": UNCLEAR},
        {"head": "Notes", "width": 28, "kind": "text"},
        {"head": "id", "width": 8, "kind": "id"}]


def make(args):
    prompts = json.loads((REPO / "data" / "prompts.json").read_text(encoding="utf-8"))
    docs = blind_check.load_docs()
    chosen = blind_check.fresh_sample(docs, N_DOCS, random.Random(SEED), sets={"reviews"}, restricted=True)
    if len(chosen) < N_DOCS:
        sys.exit(f"only {len(chosen)} fresh review documents found; is data/local present?")
    qs = []
    for i, d in enumerate(chosen, 1):
        lst = prompts["lists"][d["list_type"]]
        what = f"{lst['many'].capitalize()}: {lst['describe']}"
        qs.append({"vals": [i, d["text"], what, None, None, None, None, d["doc_id"]], "group": None})
    for person in ("A", "B"):
        wb, lists, cache = forms.xl_workbook(person)
        forms.xl_sheet(wb, lists, cache, "Reviews", COLS, qs)
        forms.xl_finish(wb, Path(args.out) / f"blind-reviews-{person}.xlsx")
    KEY.parent.mkdir(exist_ok=True)
    with KEY.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["doc_id", "n"])
        for d in chosen:
            w.writerow([d["doc_id"], d["n"]])
    print(f"wrote blind-reviews-A.xlsx and -B.xlsx in {args.out} ({N_DOCS} documents, n = "
          f"{sorted(d['n'] for d in chosen)}); key in {KEY}")


def read(path):
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True)["Reviews"]
    heads = [c.value for c in ws[1]]
    i_n, i_u, i_id = heads.index("How many items?"), heads.index(COLS[5]["head"]), heads.index("id")
    i_items, i_notes = heads.index("Items you found, one per line"), heads.index("Notes")
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[i_id]:
            out[row[i_id]] = {"n": row[i_n], "unclear": str(row[i_u] or "").lower().startswith("y"),
                              "items": row[i_items] or "", "notes": row[i_notes] or ""}
    return out


def score(args):
    key = {r["doc_id"]: int(r["n"]) for r in csv.DictReader(KEY.open(encoding="utf-8"))}
    a, b = read(args.a), read(args.b)
    blank = [(who, d) for who, s in (("A", a), ("B", b)) for d in key if not isinstance(s.get(d, {}).get("n"), int)]
    if blank:
        sys.exit(f"count missing or not a whole number: {blank}")
    for who, s in (("A", a), ("B", b)):
        print(f"author {who} matches the list on {sum(s[d]['n'] == key[d] for d in key)} of {len(key)}")
    print(f"A and B agree on {sum(a[d]['n'] == b[d]['n'] for d in key)} of {len(key)}; "
          f"unclear by either: {sum(a[d]['unclear'] or b[d]['unclear'] for d in key)}")
    for d in key:
        if a[d]["n"] != key[d] or b[d]["n"] != key[d]:
            print(f"  {d}: list {key[d]}, A {a[d]['n']}, B {b[d]['n']}")
    if args.out:
        # counts and flags only: items and notes quote review text, which stays out of git
        with open(args.out, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["doc_id", "n_list", "n_A", "n_B", "unclear_A", "unclear_B"])
            for d in key:
                w.writerow([d, key[d], a[d]["n"], b[d]["n"], a[d]["unclear"], b[d]["unclear"]])
        print(f"wrote {args.out}")


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)
    m = sub.add_parser("make")
    m.add_argument("--out", required=True)
    s = sub.add_parser("score")
    s.add_argument("a")
    s.add_argument("b")
    s.add_argument("--out", help="write counts and unclear flags (no review text) to this CSV")
    args = p.parse_args()
    {"make": make, "score": score}[args.cmd](args)


if __name__ == "__main__":
    main()
